"""Импорт учеников пачкой — текстовая вставка, CSV, XLSX.

Все методы возвращают `ImportResult`: сколько добавлено, сколько пропущено
(с причиной). Дубликаты по (name, class_label) в пределах одного класса
пропускаются, чтобы повторный импорт не создавал клонов.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from typing import TYPE_CHECKING
from uuid import UUID

from app.core.errors import Conflict

if TYPE_CHECKING:
    from app.services.students_service import StudentsService

log = logging.getLogger("mathsite.students.import")

_MAX_ROWS = 500


@dataclass(frozen=True)
class ImportRowError:
    line: int
    source: str
    reason: str


@dataclass(frozen=True)
class ImportResult:
    created: int
    skipped_duplicates: int
    errors: list[ImportRowError]


class StudentsImportService:
    def __init__(self, students: StudentsService):
        self._students = students

    # ---- Текстовая вставка ----
    async def from_text(
        self,
        text: str,
        *,
        default_class_label: str = "",
        created_by: UUID | None,
    ) -> ImportResult:
        parsed = _parse_text(text)
        if len(parsed) > _MAX_ROWS:
            raise Conflict(f"слишком много строк, максимум {_MAX_ROWS}")
        return await self._bulk_create(parsed, default_class_label, created_by)

    # ---- CSV ----
    async def from_csv(
        self,
        content: bytes,
        *,
        default_class_label: str = "",
        created_by: UUID | None,
    ) -> ImportResult:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1251", errors="replace")

        reader = csv.reader(io.StringIO(text))
        rows: list[tuple[int, str, str]] = []  # (line, name, class_label)
        first_row = True
        for i, row in enumerate(reader, start=1):
            if not row or all(not c.strip() for c in row):
                continue
            name = (row[0] if len(row) > 0 else "").strip()
            cls = (row[1] if len(row) > 1 else "").strip()
            if first_row and _looks_like_header(name, cls):
                first_row = False
                continue
            first_row = False
            if not name:
                continue
            rows.append((i, name, cls))

        if len(rows) > _MAX_ROWS:
            raise Conflict(f"слишком много строк, максимум {_MAX_ROWS}")
        return await self._bulk_create(rows, default_class_label, created_by)

    # ---- XLSX ----
    async def from_xlsx(
        self,
        content: bytes,
        *,
        default_class_label: str = "",
        created_by: UUID | None,
    ) -> ImportResult:
        try:
            from openpyxl import load_workbook
        except ImportError as e:  # pragma: no cover
            raise Conflict("openpyxl не установлен на сервере") from e

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:
            raise Conflict(f"не удалось прочитать XLSX: {e}") from e

        ws = wb.active
        if ws is None:
            raise Conflict("пустой XLSX")

        rows: list[tuple[int, str, str]] = []
        first_row = True
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if not row:
                continue
            name = (str(row[0]) if len(row) > 0 and row[0] is not None else "").strip()
            cls = (str(row[1]) if len(row) > 1 and row[1] is not None else "").strip()
            if first_row and _looks_like_header(name, cls):
                first_row = False
                continue
            first_row = False
            if not name:
                continue
            rows.append((i, name, cls))

        if len(rows) > _MAX_ROWS:
            raise Conflict(f"слишком много строк, максимум {_MAX_ROWS}")
        return await self._bulk_create(rows, default_class_label, created_by)

    # ---- Общий путь создания ----
    async def _bulk_create(
        self,
        rows: list[tuple[int, str, str]],
        default_class_label: str,
        created_by: UUID | None,
    ) -> ImportResult:
        existing = {(s.name.lower(), s.class_label.lower()) for s in await self._students.list()}

        created = 0
        skipped = 0
        errors: list[ImportRowError] = []

        for line, name, cls in rows:
            cls_label = (cls or default_class_label).strip()
            key = (name.lower(), cls_label.lower())
            if key in existing:
                skipped += 1
                continue
            try:
                await self._students.create(
                    name=name,
                    class_label=cls_label,
                    note=None,
                    created_by=created_by,
                )
                existing.add(key)
                created += 1
            except Exception as e:  # noqa: BLE001
                errors.append(
                    ImportRowError(line=line, source=f"{name} · {cls_label}", reason=str(e))
                )

        log.info("import done · created=%d skipped=%d errors=%d", created, skipped, len(errors))
        return ImportResult(created=created, skipped_duplicates=skipped, errors=errors)


# ---------- helpers ----------


def _parse_text(text: str) -> list[tuple[int, str, str]]:
    """
    Форматы строки:
        "Иванов Иван"
        "Иванов Иван, 8А"
        "Иванов Иван; 8А"
        "Иванов Иван | 8А"
        "Иванов Иван\\t8А"
    Один ученик на строке. Пустые строки игнорируются.
    """
    out: list[tuple[int, str, str]] = []
    for i, raw in enumerate(text.splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue
        parts = _SPLIT_PATTERN.split(line, maxsplit=1)
        name = parts[0].strip()
        cls = parts[1].strip() if len(parts) > 1 else ""
        if not name:
            continue
        out.append((i, name, cls))
    return out


import re

_SPLIT_PATTERN = re.compile(r"[,;\|\t]+|\s{2,}")


def _looks_like_header(a: str, b: str) -> bool:
    a_low = a.lower().strip()
    b_low = b.lower().strip()
    return a_low in {"name", "имя", "фио", "ученик"} or b_low in {
        "class",
        "класс",
        "class_label",
    }
